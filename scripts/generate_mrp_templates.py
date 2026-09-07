from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "templates"
OUT.mkdir(exist_ok=True)

BLUE = "1F4E78"
BLUE_LIGHT = "D9EAF7"
YELLOW = "FFF2CC"
GREEN = "E2F0D9"
GRAY = "E7E6E6"
WHITE = "FFFFFF"
RED = "F4CCCC"
thin = Side(style="thin", color="B7C3D0")
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def add_months(year: int, month: int, offset: int) -> date:
    index = year * 12 + month - 1 + offset
    return date(index // 12, index % 12 + 1, 1)


def base_sheet(ws, title, freeze):
    ws.title = title
    ws.freeze_panes = freeze
    ws.sheet_view.showGridLines = False
    ws.auto_filter.ref = None


def style_header(cell, fill=BLUE, color=WHITE):
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.font = Font(color=color, bold=True)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = border


def style_body(cell, locked=False, fill=None):
    cell.border = border
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    cell.protection = Protection(locked=locked)
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)


def add_instruction_sheet(wb, rows):
    ws = wb.create_sheet("填写说明", 0)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 88
    ws["A1"] = "项目"
    ws["B1"] = "说明"
    style_header(ws["A1"])
    style_header(ws["B1"])
    for idx, (item, desc) in enumerate(rows, 2):
        ws.cell(idx, 1, item)
        ws.cell(idx, 2, desc)
        style_body(ws.cell(idx, 1), fill=BLUE_LIGHT)
        style_body(ws.cell(idx, 2))
    ws.row_dimensions[1].height = 28
    for r in range(2, len(rows) + 2):
        ws.row_dimensions[r].height = 40


def build_forecast_template():
    wb = Workbook()
    ws = wb.active
    base_sheet(ws, "经营计划导入", "I3")

    add_instruction_sheet(wb, [
        ("使用流程", "先下载模板，在“经营计划导入”工作表更新数据，再上传系统。不要修改工作表名称和前两行表头结构。"),
        ("月份识别", "I:N 六个 Forecast 表头必须是 Excel 真实日期值，并设置为 yyyy年m月。系统读取单元格年月，不根据文件名、上传日期或“8月”等文本猜测年份。支持跨年，例如 2026-12 后必须是 2027-01。"),
        ("月份校验", "六个月份必须连续、不得重复、不得为空。导入前系统展示识别结果，例如：2026-08 至 2027-01，用户确认后再导入。"),
        ("料号格式", "料号列必须使用文本格式，避免 6830AA… 被转换为日期、科学计数法或丢失前导零。"),
        ("数量格式", "Forecast 允许 0 和非负数；空白表示缺失，需要修正后才能导入，不等同于 0。"),
        ("版本机制", "每次成功导入生成新的经营计划版本并保留历史；是否自动触发排产重算由排产页开关决定。"),
        ("真实样本依据", "字段来自《经营计划表.xlsx》：业务线/工厂、形态、料号、名称、项目、平台、模具、状态、连续 6 个月 fcst。"),
    ])

    headers = ["业务线/工厂", "形态", "料号", "名称", "项目", "平台", "模具", "状态"]
    for col, value in enumerate(headers, 1):
        ws.merge_cells(start_row=1, start_column=col, end_row=2, end_column=col)
        ws.cell(1, col, value)
        style_header(ws.cell(1, col))

    months = [add_months(2026, 8, i) for i in range(6)]
    for idx, month in enumerate(months, 9):
        ws.cell(1, idx, month)
        ws.cell(1, idx).number_format = 'yyyy"年"m"月"'
        ws.cell(2, idx, "fcst")
        style_header(ws.cell(1, idx), fill="4472C4")
        style_header(ws.cell(2, idx), fill="5B9BD5")
    ws.merge_cells("O1:O2")
    ws["O1"] = "fcst-total（系统校验）"
    style_header(ws["O1"], fill="548235")

    examples = [
        ["记录仪-永惠", "单机", "6830AA800561", "70迈智能记录仪M310 Pro 2K（国内版）", "BYC", "DR1400", "A01", "国内", 6260, 6480, 7030, 6930, 6730, 6730],
        ["记录仪-永惠", "单机", "6830AA800562", "70迈智能记录仪M310 Pro 3K（国内版）", "BYC", "DR1400", "A51", "国内", 6020, 6430, 6990, 6810, 6690, 6690],
    ]
    for r, row in enumerate(examples, 3):
        for c, value in enumerate(row, 1):
            ws.cell(r, c, value)
        ws.cell(r, 15, f"=SUM(I{r}:N{r})")

    for r in range(3, 503):
        for c in range(1, 16):
            fill = GREEN if c <= 14 else GRAY
            style_body(ws.cell(r, c), locked=(c == 15), fill=fill if r <= 4 else None)
        ws.cell(r, 3).number_format = "@"
        if r > 4:
            ws.cell(r, 15, f'=IF(COUNTA(I{r}:N{r})=0,"",SUM(I{r}:N{r}))')
            ws.cell(r, 15).fill = PatternFill("solid", fgColor=GRAY)

    widths = [18, 12, 18, 44, 12, 14, 10, 14] + [13] * 6 + [20]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 24
    ws.auto_filter.ref = "A2:O502"
    ws.protection.sheet = True
    ws.protection.enable()
    for row in ws.iter_rows(min_row=3, max_row=502, min_col=1, max_col=14):
        for cell in row:
            cell.protection = Protection(locked=False)

    path = OUT / "MRP经营计划导入模板.xlsx"
    wb.save(path)
    return path


def build_schedule_template():
    wb = Workbook()
    ws = wb.active
    base_sheet(ws, "12周排产计划", "H3")

    add_instruction_sheet(wb, [
        ("模板用途", "用于排产计划导出、线下核对或后续确认后的调整导入。当前需求已确认排产导出，尚未确认排产 Excel 反向导入，因此系统界面不应默认承诺反向导入。"),
        ("真实样本依据", "基础字段和布局来自《永惠-70迈计划2026年8-3.xlsx》：机型、版本、机头料号、整机料号、备注、合计。原文件虽为日排产，本项目不沿用日粒度，统一转换为连续 12 周。"),
        ("周表头", "每列使用真实周开始日期（周一）作为 Excel 日期值，显示为 yyyy-mm-dd，并在第二行显示物理周与来源月份。跨年时直接使用 2027 年真实日期。"),
        ("跨月承接", "上月第 4 周承接下月第一份；来源月份必须单独记录，不能只根据周所在月份判断。"),
        ("料号格式", "机头料号和整机料号均为文本格式，避免 Excel 自动转日期或科学计数法。"),
        ("编辑范围", "黄色周数量为可编辑示例区；灰色自动值和合计为公式或系统输出。正式导出时可按用户选择包含优先级和实际完成数量。"),
    ])

    fixed = ["机型", "版本", "机头料号", "整机料号", "备注", "工厂", "产线"]
    for c, h in enumerate(fixed, 1):
        ws.merge_cells(start_row=1, start_column=c, end_row=2, end_column=c)
        ws.cell(1, c, h)
        style_header(ws.cell(1, c))

    start = date(2026, 8, 24)
    sources = ["2026-09", "2026-09", "2026-09", "2026-09", "2026-10", "2026-10", "2026-10", "2026-10", "2026-11", "2026-11", "2026-11", "2026-11"]
    notes = ["9月提前量", "来源 9月", "来源 9月", "来源 9月", "10月提前量", "来源 10月", "来源 10月", "来源 10月", "11月提前量", "来源 11月", "来源 11月", "来源 11月"]
    from datetime import timedelta
    for i in range(12):
        c = 8 + i
        ws.cell(1, c, start + timedelta(days=7 * i))
        ws.cell(1, c).number_format = "yyyy-mm-dd"
        ws.cell(2, c, notes[i])
        style_header(ws.cell(1, c), fill="C55A11" if "提前量" in notes[i] else "4472C4")
        style_header(ws.cell(2, c), fill="F4B183" if "提前量" in notes[i] else "5B9BD5")

    tail = ["12周合计", "优先级", "实际完成数量"]
    for offset, h in enumerate(tail, 20):
        ws.merge_cells(start_row=1, start_column=offset, end_row=2, end_column=offset)
        ws.cell(1, offset, h)
        style_header(ws.cell(1, offset), fill="548235")

    examples = [
        ["BYC_DR1400_A01黑色国内版 2K版", "黑色国内版 2K版", "6810AA800175", "6830AA800561", "8月3000,9月6000,10月8000", "永惠", "一线", 2000, 1334, 1333, 1333, 2667, 1778, 1778, 1777, 0, 0, 0, 0, None, 1, 0],
        ["BYC_DR1400_A51黑色国内版 3K版", "黑色国内版 3K版", "6810AA800180", "6830AA800562", "9月4000,10月7000", "永惠", "一线", 1334, 889, 889, 888, 2334, 1556, 1555, 1555, 0, 0, 0, 0, None, 2, 0],
    ]
    for r, row in enumerate(examples, 3):
        for c, value in enumerate(row, 1):
            ws.cell(r, c, value)
        ws.cell(r, 20, f"=SUM(H{r}:S{r})")

    priority_dv = DataValidation(type="whole", operator="between", formula1="1", formula2="99", allow_blank=True)
    ws.add_data_validation(priority_dv)
    priority_dv.add("U3:U502")
    for r in range(3, 503):
        for c in range(1, 23):
            fill = YELLOW if 8 <= c <= 19 else (GRAY if c == 20 else None)
            style_body(ws.cell(r, c), locked=(c == 20), fill=fill if r <= 4 else None)
        ws.cell(r, 3).number_format = "@"
        ws.cell(r, 4).number_format = "@"
        if r > 4:
            ws.cell(r, 20, f'=IF(COUNTA(H{r}:S{r})=0,"",SUM(H{r}:S{r}))')
            ws.cell(r, 20).fill = PatternFill("solid", fgColor=GRAY)

    widths = [34, 24, 20, 20, 34, 12, 12] + [14] * 12 + [14, 10, 16]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 26
    ws.auto_filter.ref = "A2:V502"
    ws.protection.sheet = True
    ws.protection.enable()
    for row in ws.iter_rows(min_row=3, max_row=502, min_col=1, max_col=19):
        for cell in row:
            cell.protection = Protection(locked=False)
    for row in ws.iter_rows(min_row=3, max_row=502, min_col=21, max_col=22):
        for cell in row:
            cell.protection = Protection(locked=False)

    path = OUT / "MRP十二周排产计划模板.xlsx"
    wb.save(path)
    return path


def build_inventory_template():
    wb = Workbook()
    ws = wb.active
    base_sheet(ws, "库存快照导入", "A2")

    add_instruction_sheet(wb, [
        ("使用流程", "先下载模板，在“库存快照导入”工作表更新数据，再上传系统。不要修改工作表名称和第一行表头。"),
        ("库存口径", "永惠和 APK 分开计算；同一工厂有效仓库汇总；不良品仓不计入月可排量。"),
        ("料号格式", "物料编码必须使用文本格式，避免被转换为日期、科学计数法或丢失前导零。"),
        ("快照日期", "使用 Excel 真实日期值，表示该行库存所属的统一快照日期。"),
        ("数量格式", "数量必须为 0 或非负数；空白表示缺失，导入前需要修正。"),
        ("自动重算", "导入成功后是否自动触发排产重算，由排产页面的库存自动重算开关决定。"),
    ])

    headers = ["物料编码", "仓库编码", "仓库名称", "仓库类型", "工厂编码", "数量", "快照日期", "是否计入计算"]
    for c, header in enumerate(headers, 1):
        ws.cell(1, c, header)
        style_header(ws.cell(1, c))

    examples = [
        ["6830AA800561", "YH-FG-01", "永惠成品仓", "良品仓", "永惠", 200, date(2026, 8, 28), "是"],
        ["6830AA800561", "YH-BAD-01", "永惠不良品仓", "不良品仓", "永惠", 46, date(2026, 8, 28), "否"],
    ]
    for r, row in enumerate(examples, 2):
        for c, value in enumerate(row, 1):
            ws.cell(r, c, value)

    include_dv = DataValidation(type="list", formula1='"是,否"', allow_blank=False)
    ws.add_data_validation(include_dv)
    include_dv.add("H2:H502")
    for r in range(2, 503):
        for c in range(1, 9):
            style_body(ws.cell(r, c), locked=False, fill=GREEN if r <= 3 else None)
        ws.cell(r, 1).number_format = "@"
        ws.cell(r, 7).number_format = "yyyy-mm-dd"

    widths = [20, 18, 24, 14, 14, 12, 14, 16]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 30
    ws.auto_filter.ref = "A1:H502"
    path = OUT / "MRP库存快照导入模板.xlsx"
    wb.save(path)
    return path


if __name__ == "__main__":
    for output in (build_forecast_template(), build_schedule_template(), build_inventory_template()):
        print(output)
